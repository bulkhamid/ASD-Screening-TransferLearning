from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports"
ASSET_DIR = OUT_DIR / "assets"
DOCX_PATH = OUT_DIR / "ASD_Project_Comprehensive_Update_Gardos.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_GRAY = "F2F4F7"
CALL_OUT = "F4F6F9"
RISK_RED = "9B1C1C"
GOLD = "7A5A00"
GREEN = "2E7D32"
INK = "1A1A1A"


def ensure_dirs():
    OUT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)


def font(size=18, bold=False):
    candidates = [
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def load_json(path):
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


def draw_bar_chart(title, rows, labels, colors, path, width=1200, height=620):
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    title_font = font(30, True)
    label_font = font(18)
    small_font = font(16)
    d.text((40, 30), title, fill="#" + INK, font=title_font)
    left, top, right, bottom = 330, 110, width - 70, height - 110
    d.line((left, bottom, right, bottom), fill="#444444", width=2)
    d.line((left, top, left, bottom), fill="#444444", width=2)
    for i in range(6):
        x = left + i * (right - left) / 5
        d.line((x, bottom, x, bottom + 8), fill="#444444", width=1)
        d.text((x - 10, bottom + 12), f"{i/5:.1f}", fill="#444444", font=small_font)
        d.line((x, top, x, bottom), fill="#E8E8E8", width=1)
    group_h = (bottom - top) / max(len(rows), 1)
    bar_h = min(24, group_h / (len(labels) + 0.8))
    for idx, row in enumerate(rows):
        y_mid = top + group_h * idx + group_h / 2
        label = row["name"]
        if len(label) > 28:
            words = label.split()
            first = []
            second = []
            for word in words:
                if len(" ".join(first + [word])) <= 24:
                    first.append(word)
                else:
                    second.append(word)
            label = " ".join(first) + "\n" + " ".join(second)
            d.multiline_text((40, y_mid - 24), label, fill="#222222", font=label_font, spacing=2)
        else:
            d.text((40, y_mid - 16), label, fill="#222222", font=label_font)
        for j, label in enumerate(labels):
            value = row[label]
            y = y_mid - (len(labels) * bar_h) / 2 + j * bar_h
            x2 = left + value * (right - left)
            d.rectangle((left, y, x2, y + bar_h * 0.72), fill=colors[j])
            d.text((x2 + 8, y - 1), f"{value:.3f}", fill="#222222", font=small_font)
    legend_x = left
    for j, label in enumerate(labels):
        y = height - 55
        x = legend_x + j * 210
        d.rectangle((x, y, x + 22, y + 14), fill=colors[j])
        d.text((x + 30, y - 3), label.replace("_", " ").title(), fill="#222222", font=small_font)
    img.save(path, quality=95)


def draw_grouped_counts(title, groups, path, width=1200, height=650):
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    d.text((40, 30), title, fill="#" + INK, font=font(30, True))
    left, top, right, bottom = 120, 120, width - 80, height - 110
    max_val = max(max(v["positive"], v["negative"]) for v in groups.values())
    max_val = max(max_val, 1)
    d.line((left, bottom, right, bottom), fill="#444444", width=2)
    bar_w = (right - left) / max(len(groups), 1) / 4
    for idx, (name, counts) in enumerate(groups.items()):
        cx = left + (idx + 0.5) * (right - left) / len(groups)
        for j, key in enumerate(["positive", "negative"]):
            val = counts[key]
            x1 = cx + (j - 1) * bar_w
            x2 = x1 + bar_w * 0.85
            y1 = bottom - (val / max_val) * (bottom - top)
            color = "#2E74B5" if key == "positive" else "#7A869A"
            d.rectangle((x1, y1, x2, bottom), fill=color)
            d.text((x1, y1 - 22), str(val), fill="#222222", font=font(15))
        d.text((cx - 35, bottom + 16), name, fill="#222222", font=font(16))
    d.rectangle((left, height - 62, left + 22, height - 48), fill="#2E74B5")
    d.text((left + 30, height - 66), "Positive", fill="#222222", font=font(16))
    d.rectangle((left + 150, height - 62, left + 172, height - 48), fill="#7A869A")
    d.text((left + 180, height - 66), "Negative", fill="#222222", font=font(16))
    img.save(path, quality=95)


def draw_effect_chart(title, effects, path, width=1200, height=660):
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    d.text((40, 30), title, fill="#" + INK, font=font(30, True))
    left, mid, right = 260, 640, width - 80
    top, bottom = 110, height - 80
    d.line((mid, top, mid, bottom), fill="#444444", width=2)
    d.text((mid - 20, bottom + 10), "0", fill="#444444", font=font(15))
    max_abs = max(abs(x["d"]) for x in effects) or 1
    row_h = (bottom - top) / len(effects)
    for i, item in enumerate(effects):
        y = top + i * row_h + row_h / 2
        d.text((40, y - 12), item["feature"], fill="#222222", font=font(17))
        val = item["d"]
        x = mid + (val / max_abs) * (right - mid if val > 0 else mid - left)
        color = "#2E74B5" if val > 0 else "#9B1C1C"
        d.rectangle((min(mid, x), y - 11, max(mid, x), y + 11), fill=color)
        d.text((x + (8 if val > 0 else -58), y - 12), f"{val:.2f}", fill="#222222", font=font(15))
    d.text((left, height - 45), "Negative-associated", fill="#9B1C1C", font=font(15, True))
    d.text((right - 170, height - 45), "Positive-associated", fill="#2E74B5", font=font(15, True))
    img.save(path, quality=95)


def make_gradcam_grid(path):
    meta = load_json("gradcam_outputs/inspection_app/metadata.json")
    samples = meta["samples"]
    selected = [samples[0], next(s for s in samples if s["label"] == 0)]
    cells = []
    for sample in selected:
        frame = sample["frames"][0]
        base = ROOT / "gradcam_outputs" / "inspection_app"
        label = "ASD-positive" if sample["label"] else "Control"
        cells.append((label + " original", base / frame["original"]))
        cells.append((label + " R3D Grad-CAM", base / frame["r3d"]["image"]))
        cells.append((label + " Simple3DLSTM Grad-CAM", base / frame["simple3dlstm"]["image"]))

    thumb_w, thumb_h = 300, 300
    pad = 26
    title_h = 38
    img = Image.new("RGB", (3 * thumb_w + 4 * pad, 2 * (thumb_h + title_h) + 3 * pad), "white")
    d = ImageDraw.Draw(img)
    for idx, (label, src) in enumerate(cells):
        row, col = divmod(idx, 3)
        x = pad + col * (thumb_w + pad)
        y = pad + row * (thumb_h + title_h + pad)
        d.text((x, y), label, fill="#222222", font=font(18, True))
        im = Image.open(src).convert("RGB").resize((thumb_w, thumb_h))
        img.paste(im, (x, y + title_h))
        d.rectangle((x, y + title_h, x + thumb_w, y + title_h + thumb_h), outline="#D0D5DD", width=2)
    img.save(path, quality=95)


def setup_doc_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=INK):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, h in enumerate(headers):
        shade_cell(table.rows[0].cells[i], LIGHT_GRAY)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_callout(doc, title, body, fill=CALL_OUT):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    r.font.size = Pt(11)
    p.add_run("\n" + body)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_figure(doc, path, caption, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("title", caption[:120])
        drawing[0].set("descr", caption[:250])
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)


def generate_assets():
    shortcut_source = load_json("experiments/shortcut_baselines_source_aware_results.json")["results"]
    rows = [
        {"name": r["name"].replace("_", " "), "test_auc": r["test"]["auc"], "test_f1": r["test"]["f1"]}
        for r in shortcut_source
    ]
    draw_bar_chart(
        "Shortcut Baselines on Source-Aware Cache",
        rows,
        ["test_auc", "test_f1"],
        ["#2E74B5", "#7A869A"],
        ASSET_DIR / "shortcut_source_aware_chart.jpg",
    )

    model_rows = [
        {"name": "Simple2DCNN", "test_auc": 1.000, "test_f1": 1.000},
        {"name": "Simple3DLSTM", "test_auc": 0.955, "test_f1": 0.952},
        {"name": "R3D fine-tuned", "test_auc": 1.000, "test_f1": 1.000},
    ]
    draw_bar_chart(
        "50-Epoch Source-Aware Model Results",
        model_rows,
        ["test_auc", "test_f1"],
        ["#2E74B5", "#2E7D32"],
        ASSET_DIR / "model_source_aware_chart.jpg",
    )

    diag = load_json("data_source_aware/split_diagnostics.json")
    draw_grouped_counts(
        "Overall Capture Month by Label",
        diag["overall"]["month_counts"],
        ASSET_DIR / "month_counts_chart.jpg",
    )

    feat = load_json("experiments/feature_distribution_source_aware_results.json")
    effects = [
        {"feature": r["feature"], "d": r["cohen_d_pos_minus_neg"]}
        for r in feat["feature_summary"]["test"]["features"][:8]
    ]
    draw_effect_chart(
        "Top Test-Set Feature Differences (Cohen's d)",
        effects,
        ASSET_DIR / "feature_effects_chart.jpg",
    )

    make_gradcam_grid(ASSET_DIR / "gradcam_grid.jpg")


def build_doc():
    doc = Document()
    setup_doc_styles(doc)
    header = doc.sections[0].header.paragraphs[0]
    header.text = "ASD Screening Project Technical Update"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string("667085")
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Prepared for Professor Gardos"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor.from_string("667085")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ASD Screening Project: Professor Feedback Response, Audit Findings, and Next-Step Evaluation Plan")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Comprehensive technical update for Professor Gardos").italic = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Prepared by Zhamshidbek Abdulkhamidov | June 2026")

    add_callout(
        doc,
        "Bottom line",
        "The professor's main implementation suggestions were addressed, and the models still perform strongly. However, the strongest new finding is that simple non-neural shortcut baselines also perform very well, so the current results should be framed as promising but not yet evidence of robust behavioral ASD generalization.",
        fill="E8EEF5",
    )

    doc.add_heading("1. Executive Summary", level=1)
    add_bullets(doc, [
        "Aspect-ratio-preserving letterbox preprocessing is implemented with gray padding, replacing direct 112x112 stretching.",
        "BCEWithLogitsLoss is now documented correctly as sigmoid plus binary cross-entropy in one numerically stable loss.",
        "R3D supports an optional full-backbone fine-tuning stage; the 50-epoch source-aware run used the requested two-stage head plus layer4/fc plan.",
        "Cached --cv is blocked because the old cached CV repeated the same train/val/test split and was not valid 5-fold group CV.",
        "Grad-CAM tooling and dashboards were added; padding is masked so heatmaps do not visually reward letterbox borders.",
        "The source-aware split improved the audit discipline but did not remove dataset shortcuts: motion statistics, thumbnails, and brightness/color features remain highly predictive.",
    ])

    doc.add_heading("2. Professor Feedback: Implementation Status", level=1)
    add_table(doc, ["Suggestion", "Implementation", "Status"], [
        ["Preserve aspect ratio when resizing", "Frames are resized by scaling the longer side to 112 and padding the shorter side with gray value 128.", "Done"],
        ["Clarify BCEWithLogitsLoss", "Paper/docs now state that BCEWithLogitsLoss combines sigmoid activation and BCE for raw single-logit outputs.", "Done"],
        ["Add R3D full-backbone stage", "Optional flags --full_finetune, --full_ft_epochs, and --full_ft_lr were added.", "Done"],
        ["Avoid misleading cached CV", "train.py now fails fast for --cached --cv until fold-specific cached splits exist.", "Done"],
        ["Inspect model behavior", "Grad-CAM dashboards show R3D layer4 and Simple3DLSTM cnn.10; layer-by-layer dashboards cover train/val/test.", "Done"],
        ["Check simple baselines", "Simple2DCNN plus non-neural shortcut baselines were added; results show strong static/source cues.", "Done"],
    ], widths=[2.2, 3.4, 0.8])

    doc.add_heading("3. Data and Preprocessing Changes", level=1)
    add_bullets(doc, [
        "The clean cache contains 107 usable videos after excluding one unreadable positive video.",
        "Cached tensor shape is (N, 60, 112, 112, 3), uint8, with no all-zero videos.",
        "Both original and source-aware caches have no exact decoded-array duplicates across train/val/test.",
        "Split manifests include path, label, raw hash, decoded hash, frame counts, read status, resize mode, pad value, source month, and extension.",
    ])
    add_table(doc, ["Cache", "Train", "Validation", "Test", "Integrity"], [
        ["Gray holdout", "74 videos, 37/37", "11 videos, 5/6", "22 videos, 11/11", "No all-zero videos; no exact decoded duplicates"],
        ["Source-aware", "74 videos, 37/37", "11 videos, 5/6", "22 videos, 11/11", "No all-zero videos; no exact decoded duplicates"],
    ], widths=[1.3, 1.4, 1.4, 1.4, 1.9])

    doc.add_heading("4. Model Results", level=1)
    add_figure(doc, ASSET_DIR / "model_source_aware_chart.jpg", "Figure 1. Source-aware 50-epoch model performance. AUC is threshold-free; F1 uses the validation-selected threshold.")
    add_table(doc, ["Model", "Temp.", "Threshold", "AUC", "Acc.", "Precision", "Recall", "F1", "Loss"], [
        ["Simple2DCNN", "0.488", "0.55", "1.000", "1.000", "1.000", "1.000", "1.000", "0.0053"],
        ["Simple3DLSTM", "0.002", "0.05", "0.955", "0.955", "1.000", "0.909", "0.952", "5.4924"],
        ["R3D fine-tuned", "0.546", "0.05", "1.000", "1.000", "1.000", "1.000", "1.000", "0.0594"],
    ], widths=[1.35, 0.55, 0.7, 0.55, 0.55, 0.75, 0.65, 0.55, 0.65])
    add_callout(
        doc,
        "Calibration caution",
        "Simple3DLSTM selected a near-zero temperature (0.002), which means its calibrated probabilities should not be interpreted as reliable. Its ranking remained strong, but probability values are unstable with the current small validation set.",
        fill="FFF8E5",
    )

    doc.add_heading("5. Shortcut and Source-Bias Audit", level=1)
    add_figure(doc, ASSET_DIR / "shortcut_source_aware_chart.jpg", "Figure 2. Non-neural shortcut baselines on the source-aware cache. Strong results here indicate possible source/recording artifacts.")
    add_table(doc, ["Shortcut baseline", "Val AUC", "Test AUC", "Test Acc.", "Test F1", "Interpretation"], [
        ["Majority prior", "0.500", "0.500", "0.500", "0.667", "Chance-style baseline"],
        ["Global color/brightness/padding", "0.867", "0.901", "0.864", "0.870", "Low-level visual statistics remain predictive"],
        ["Frame-difference motion stats", "1.000", "1.000", "0.955", "0.952", "Motion/recording differences are nearly sufficient"],
        ["Tiny static thumbnails", "1.000", "0.967", "0.955", "0.957", "Static frame/background/framing cues are strong"],
        ["Path suffix/date tokens", "0.500", "0.591", "0.591", "0.308", "Nearer chance after excluding label-folder leakage"],
    ], widths=[1.8, 0.55, 0.55, 0.6, 0.55, 2.35])
    add_figure(doc, ASSET_DIR / "month_counts_chart.jpg", "Figure 3. Capture month is class-confounded in the available data.")
    add_figure(doc, ASSET_DIR / "feature_effects_chart.jpg", "Figure 4. The largest class differences are frame-difference motion statistics and color/brightness features.")
    add_callout(
        doc,
        "Scientific interpretation",
        "The source-aware split can reduce split randomness but cannot remove bias that exists in the raw collection. August/September videos are all positive, .mp4 files are all positive, and July/October are mostly negative. This is why high deep-model scores must be reported cautiously.",
        fill="FDECEC",
    )

    doc.add_heading("6. Grad-CAM Inspection", level=1)
    add_figure(doc, ASSET_DIR / "gradcam_grid.jpg", "Figure 5. Representative Grad-CAM overlays from the inspection dashboard: original frame, R3D heatmap, and Simple3DLSTM heatmap for one positive and one control example.")
    add_bullets(doc, [
        "R3D Grad-CAM is generated from net.layer4, the final convolutional block.",
        "Simple3DLSTM Grad-CAM is generated from cnn.10, the final CNN block before temporal modeling.",
        "The first black-padding issue was corrected by switching to gray padding and masking padding in overlay rendering.",
        "Grad-CAM highlights regions contributing to a decision for a chosen layer; it is not causal proof that the model learned ASD behavior.",
        "Current Grad-CAM examples look more focused on person/background content than padding, but the shortcut audit shows this visual inspection is not sufficient by itself.",
    ])

    doc.add_heading("7. Testing and Validation Strategy", level=1)
    add_table(doc, ["Test category", "Current coverage", "Purpose"], [
        ["Unit tests", "31 unittest cases", "Verify metrics, preprocessing, cached loader paths, source-aware splitting, Grad-CAM helpers, and CLI guard behavior"],
        ["CLI smoke tests", "train.py, save_numpy_dataset.py, shortcut audits, feature audit", "Ensure main user-facing commands run"],
        ["Data integrity", "Shape, class balance, all-zero detection, exact duplicate hashes", "Prevent obvious cache or split corruption"],
        ["Sanity controls", "Corrected label-shuffle checks", "Check whether models collapse toward chance when labels are randomized"],
        ["Shortcut baselines", "Brightness/color, motion stats, thumbnails, path/date tokens", "Detect whether simple artifacts can solve the split"],
        ["Interpretability", "Grad-CAM dashboards and layer inspection", "Inspect attention qualitatively without treating it as proof"],
    ], widths=[1.3, 2.2, 3.0])
    doc.add_heading("Recommended Next Validation Steps", level=2)
    add_bullets(doc, [
        "Collect or reconstruct subject IDs, session IDs, device/site/source IDs, and capture batches.",
        "Create true subject-level or source/session-level splits where one child/source never appears in multiple splits.",
        "Keep a final untouched test set; choose thresholds, temperatures, and hyperparameters only on validation data.",
        "Report AUC as the primary metric; present thresholded metrics as secondary.",
        "Run shortcut baselines beside every deep model result.",
        "Seek an external test set from a different source or collection session before making strong generalization claims.",
    ])

    doc.add_heading("8. Manuscript Revisions Needed", level=1)
    add_bullets(doc, [
        "Replace claims of rigorous cached 5-fold GroupKFold CV with fixed holdout or future fold-specific CV language.",
        "Revise the abstract and conclusion so perfect scores are described as preliminary holdout results with shortcut-bias caveats.",
        "Update preprocessing text and code listing from direct cv2.resize to aspect-ratio-preserving letterbox resize.",
        "Update R3D pretrained weight wording from deprecated weights=True to the explicit torchvision weight enum.",
        "Add the Simple2DCNN and non-neural shortcut baselines as required controls.",
        "Add the source-aware split diagnostics and explain why subject/external validation remains necessary.",
    ])

    doc.add_heading("9. Follow-Up Email Draft", level=1)
    email = (
        "Dear Professor Gardos,\n\n"
        "I hope you are doing well. I apologize for the late follow-up. I had classes after our discussion, and I also recently got married, so things were quite busy. I am really sorry for the delay.\n\n"
        "I wanted to follow up on the ASD screening project and the suggestions you gave me. I implemented the aspect-ratio-preserving resize, clarified the BCEWithLogitsLoss explanation, added the optional third R3D fine-tuning stage, and corrected the evaluation language so the cached repeated split is no longer described as valid 5-fold cross-validation.\n\n"
        "I also added Grad-CAM inspection, a source-aware holdout split, and several shortcut baselines. The important new finding is that the deep models still perform strongly, but simple non-neural features such as motion statistics, brightness/color, and static thumbnails also classify the current split very well. This suggests that the current dataset likely contains source or recording artifacts, so I am treating the perfect results cautiously and focusing next on stronger subject/source-independent validation.\n\n"
        "I am planning to apply for PhD programs for the 2027 cycle, and this research project would be very important for strengthening my application. I would really appreciate your feedback on the attached technical update and on how to make the evaluation more rigorous before I present or submit the work.\n\n"
        "Thank you again for your guidance.\n\n"
        "Best regards,\n"
        "Zhamshidbek Abdulkhamidov"
    )
    for para in email.split("\n\n"):
        p = doc.add_paragraph(para)
        p.paragraph_format.space_after = Pt(8)

    doc.add_heading("10. Appendix: Commands and Artifacts", level=1)
    add_bullets(doc, [
        "Source-aware cache: data_source_aware/",
        "Diagnostics: data_source_aware/split_diagnostics.json",
        "Shortcut results: experiments/shortcut_baselines_source_aware_results.json",
        "Feature audit: experiments/feature_distribution_source_aware_results.json",
        "Evaluation protocol: docs/EvaluationProtocol.md",
        "Audit report: docs/AuditReport.md",
        "GitHub branch pushed: Zhamshidbek-updates at commit 6a08d51.",
    ])

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    ensure_dirs()
    generate_assets()
    build_doc()
    print(DOCX_PATH)
